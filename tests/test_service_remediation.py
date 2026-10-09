"""Regression contracts for identity, snapshots, Ask, extraction and BYOC.

All model-provider calls and embeddings in these isolated corpora are local
fakes. No provider key or network request is needed.
"""

import io
import json
import re
import subprocess
import sys
import threading
from pathlib import Path

import numpy as np
import pytest
from fastapi.testclient import TestClient

from src import ask, auth, embedder, indexer, ingest, main, workspaces
from src.storage import corpus_lock, corpus_transaction


def fake_embed(texts, **kwargs):
    vectors = np.array([[1.0, len(text) % 3 / 10, 0.1, 0.2] for text in texts], dtype=np.float32)
    return vectors / np.linalg.norm(vectors, axis=1, keepdims=True)


@pytest.fixture
def local_corpus(tmp_path, monkeypatch):
    corpora = tmp_path / "corpora"
    root = corpora / "own"
    (root / "documents").mkdir(parents=True)
    monkeypatch.setattr(workspaces, "CORPORA_DIR", str(corpora))
    monkeypatch.setattr(workspaces, "ARTIFACTS_ROOT", str(tmp_path / "artifacts"))
    monkeypatch.setenv("QUERYTRACE_DEFAULT_WORKSPACE", "own")
    monkeypatch.setattr(indexer, "_embed_texts", fake_embed)
    monkeypatch.setattr(embedder, "embed", fake_embed)
    (root / "workspace.json").write_text(
        json.dumps({"name": "Private research", "doc_types": ["document"]})
    )
    (root / "roles.json").write_text(
        json.dumps({"roles": {"viewer": {"access_rank": 1}, "owner": {"access_rank": 2}}})
    )
    content = "Research on annual revenue retention and financial planning. " * 5
    entry = {
        "id": "doc_001",
        "file_name": "base.txt",
        "title": "Revenue plan",
        "type": "document",
        "date": "2024-01-01",
        "min_role": "viewer",
        "sensitivity": "low",
        "superseded_by": None,
        "tags": [],
        "short_summary": content[:100],
    }
    (root / "metadata.json").write_text(json.dumps({"documents": [entry]}))
    (root / "documents" / "base.txt").write_text(content)
    workspaces.invalidate_workspace_cache()
    ask.reset_ask_state()
    ws = workspaces.get_workspace()
    indexer.build_and_save(ws)
    yield ws
    workspaces.invalidate_workspace_cache()
    ask.reset_ask_state()


def upload(ws):
    return ingest.ingest_document(
        file_bytes=b"New acquisition strategy and annual revenue retention research. " * 8,
        title="Acquisition strategy",
        date="2024-05-01",
        min_role="viewer",
        doc_type="document",
        sensitivity="low",
        tags=[],
        workspace=ws,
        filename="new.txt",
    )


@pytest.fixture
def fake_model(monkeypatch):
    calls = []
    monkeypatch.setenv("ANTHROPIC_API_KEY", "dummy-no-network")
    monkeypatch.setenv("ASK_MAX_PER_MINUTE", "100")

    def model(system, user):
        calls.append(user)
        ids = re.findall(r'<doc id="([^"]+)"', user)
        return ask.ModelReply("Claim [" + ", ".join(ids + ["doc_999"]) + "].", "mock", 777, 33)

    monkeypatch.setattr(ask, "call_model", model)
    return calls


def test_custom_workspace_without_optional_files_works(local_corpus):
    client = TestClient(main.app)
    meta = client.get("/workspaces/own/meta").json()
    assert meta["personas"] == [] and meta["scenarios"] == []
    assert meta["is_demo"] is False and meta["evaluation_available"] is False
    assert client.get("/documents?workspace=own").json()["visible_docs"] == 1
    response = client.get("/evals?workspace=own")
    assert response.status_code == 200
    assert response.json()["status"] == "unavailable"
    assert response.json()["aggregate"]["queries_run"] == 0


def test_non_demo_user_hints_never_exposed(local_corpus):
    Path(local_corpus.users_path).write_text(
        json.dumps(
            {
                "users": [
                    {
                        "username": "private",
                        "name": "Private",
                        "role": "viewer",
                        "password_hint": "not-public",
                        "password_sha256": auth.hash_password("secret"),
                    }
                ]
            }
        )
    )
    client = TestClient(main.app)
    assert client.get("/personas?workspace=own").json()["personas"] == []
    assert client.get("/workspaces/own/meta").json()["personas"] == []


def test_disabled_lab_rejects_guest_documents_and_queries(local_corpus):
    Path(local_corpus.root, "workspace.json").write_text(
        json.dumps({"name": "Private", "lab_enabled": False})
    )
    client = TestClient(main.app)
    assert client.get("/documents?workspace=own").status_code == 401
    assert (
        client.post(
            "/query", json={"query": "annual revenue", "workspace": "own", "role": "viewer"}
        ).status_code
        == 401
    )


def test_invalid_default_is_actionable(local_corpus, monkeypatch):
    monkeypatch.setenv("QUERYTRACE_DEFAULT_WORKSPACE", "missing")
    with pytest.raises(
        workspaces.WorkspaceConfigurationError, match="QUERYTRACE_DEFAULT_WORKSPACE"
    ):
        workspaces.validate_startup()


def test_invalid_role_configuration_is_actionable(local_corpus):
    Path(local_corpus.roles_path).write_text('{"roles":{"viewer":{"access_rank":"high"}}}')
    with pytest.raises(workspaces.WorkspaceConfigurationError, match="access_rank"):
        workspaces.get_workspace("own")


@pytest.mark.parametrize("stage", ["metadata", "bm25", "payloads", "faiss"])
def test_publication_failure_restores_all_files_and_next_upload_succeeds(
    local_corpus, monkeypatch, stage
):
    ws = local_corpus
    paths = [ws.metadata_path, ws.index_path, ws.index_docs_path, ws.bm25_path]
    before = {p: Path(p).read_bytes() for p in paths}
    before_version = workspaces.corpus_version(ws)
    original = indexer._atomic_write_json
    original_faiss = indexer.faiss.write_index
    target = {
        "metadata": ws.metadata_path,
        "bm25": ws.bm25_path,
        "payloads": ws.index_docs_path,
    }.get(stage)

    def fail_json(path, obj, **kwargs):
        original(path, obj, **kwargs)
        if path == target:
            raise RuntimeError("injected after " + stage)

    def fail_faiss(index, path):
        original_faiss(index, path)
        raise RuntimeError("injected faiss publication")

    with monkeypatch.context() as patch:
        patch.setattr(indexer, "_atomic_write_json", fail_json)
        if stage == "faiss":
            patch.setattr(indexer.faiss, "write_index", fail_faiss)
        with pytest.raises(RuntimeError, match="injected"):
            upload(ws)
    assert {p: Path(p).read_bytes() for p in paths} == before
    assert len(list(Path(ws.documents_dir).glob("*.txt"))) == 1
    assert not Path(ws.root, ".querytrace-transaction").exists()
    assert workspaces.corpus_version(workspaces.get_workspace("own")) == before_version
    entry = upload(workspaces.get_workspace("own"))
    assert entry["id"] == "doc_002"
    idx, payloads = indexer.load_persisted_index("own")
    bm25 = indexer.load_bm25_corpus("own")
    assert idx.ntotal == len(payloads) == len(bm25)
    assert len({p["id"] for p in payloads}) == len(payloads)
    assert {p["doc_id"] for p in payloads} == {"doc_001", "doc_002"}


def test_interrupted_process_journal_recovers_before_next_read(tmp_path):
    root = tmp_path / "corpus"
    root.mkdir()
    old = root / "metadata.json"
    new = root / "document.txt"
    old.write_text("original")
    script = """import os,sys
from src.storage import corpus_lock,corpus_transaction
from pathlib import Path
root,old,new=sys.argv[1:]
with corpus_lock(root):
    with corpus_transaction(root,[old,new]):
        Path(old).write_text('partial')
        Path(new).write_text('orphan')
        os._exit(23)
"""
    result = subprocess.run([sys.executable, "-c", script, str(root), str(old), str(new)])
    assert result.returncode == 23 and old.read_text() == "partial"
    with corpus_lock(str(root)):
        assert old.read_text() == "original" and not new.exists()
    assert not (root / ".querytrace-transaction").exists()


def test_reader_waits_for_complete_transaction(tmp_path):
    path = tmp_path / "data"
    path.write_text("old")
    started, release, read_done = threading.Event(), threading.Event(), threading.Event()
    seen = []

    def writer():
        with corpus_lock(str(tmp_path)):
            with corpus_transaction(str(tmp_path), [str(path)]):
                path.write_text("partial")
                started.set()
                assert release.wait(3)
                path.write_text("complete")

    def reader():
        with corpus_lock(str(tmp_path)):
            seen.append(path.read_text())
        read_done.set()

    first = threading.Thread(target=writer)
    second = threading.Thread(target=reader)
    first.start()
    assert started.wait(3)
    second.start()
    assert not read_done.wait(0.05)
    release.set()
    first.join(3)
    second.join(3)
    assert seen == ["complete"]


def test_ask_cache_invalidated_by_successful_corpus_publication(local_corpus, fake_model):
    client = TestClient(main.app)
    body = {"query": "annual revenue retention", "role": "viewer", "policy_name": "naive_top_k"}
    before = client.post("/ask", json=body).json()
    cached = client.post("/ask", json=body).json()
    assert cached["cached"] is True and len(fake_model) == 1
    upload(local_corpus)
    after = client.post("/ask", json=body).json()
    assert after["cached"] is False and len(fake_model) == 2
    assert before["execution"]["corpus_version"] != after["execution"]["corpus_version"]
    assert "doc_002" in {d["doc_id"] for d in after["context"]}


def test_ask_context_trace_citations_exactly_match_trimmed_prompt(
    local_corpus, fake_model, monkeypatch
):
    upload(local_corpus)
    client = TestClient(main.app)
    body = {"query": "annual revenue retention", "role": "viewer", "policy_name": "naive_top_k"}
    complete = client.post("/ask", json=body).json()
    assert complete["context_trimmed"] is False
    first_tokens = complete["decision_trace"]["included"][0]["token_count"]
    monkeypatch.setenv("ASK_MAX_CONTEXT_TOKENS", str(first_tokens))
    trimmed = client.post("/ask", json=body).json()
    assert trimmed["cached"] is False and trimmed["context_trimmed"] is True
    assert len(trimmed["context"]) == 1 < len(trimmed["selected_context"])
    assert trimmed["total_tokens"] == first_tokens == trimmed["decision_trace"]["total_tokens"]
    trace = trimmed["decision_trace"]
    assert trace["metrics"]["included_count"] == 1
    assert trace["metrics"]["prompt_excluded_count"] == len(trimmed["prompt_excluded"])
    assert {d["doc_id"] for d in trace["included"]} == {d["doc_id"] for d in trimmed["context"]}
    assert set(re.findall(r'<doc id="([^"]+)"', fake_model[-1])) == {
        d["doc_id"] for d in trimmed["context"]
    }
    assert (
        trimmed["usage"]["input_tokens"] == 777
    )  # Provider usage is distinct from local estimates.
    assert next(c for c in trimmed["citations"] if c["doc_id"] == "doc_999")["valid"] is False


def test_ask_zero_prompt_budget_keeps_trace_without_provider_call(
    local_corpus, fake_model, monkeypatch
):
    monkeypatch.setenv("ASK_MAX_CONTEXT_TOKENS", "0")
    response = (
        TestClient(main.app).post("/ask", json={"query": "annual revenue", "role": "viewer"}).json()
    )
    assert response["model"] == "not_called" and fake_model == []
    assert response["context"] == [] and response["decision_trace"] is not None
    assert response["usage"]["input_tokens"] == response["total_tokens"] == 0


def test_grouped_citations_validate_each_id():
    citations = ask.extract_citations(
        "Claim [doc_001, doc_999, doc_002]. Again [doc_002].", ["doc_001", "doc_002"]
    )
    assert [(c.doc_id, c.valid) for c in citations] == [
        ("doc_001", True),
        ("doc_999", False),
        ("doc_002", True),
    ]


def test_docx_preserves_body_table_order_nested_cells_header_footer():
    import docx

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Header provenance"
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "$300 million"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested note"
    document.add_paragraph("After table")
    document.sections[0].footer.paragraphs[0].text = "Footer classification"
    stream = io.BytesIO()
    document.save(stream)
    text = ingest.extract_text(stream.getvalue(), "table.docx")
    ordered = [
        "Header provenance",
        "Before table",
        "Revenue",
        "$300 million",
        "Nested note",
        "After table",
        "Footer classification",
    ]
    assert [text.index(part) for part in ordered] == sorted(text.index(part) for part in ordered)


def test_docx_unsupported_substantive_element_rejected():
    import docx
    from docx.oxml import OxmlElement

    document = docx.Document()
    paragraph = document.add_paragraph("Financial note " * 20)
    paragraph._p.append(OxmlElement("w:footnoteReference"))
    stream = io.BytesIO()
    document.save(stream)
    with pytest.raises(ingest.IngestError, match="footnotes") as error:
        ingest.extract_text(stream.getvalue(), "footnotes.docx")
    assert error.value.status_code == 422


@pytest.mark.parametrize("hostname", ["localhost", "127.0.0.1"])
def test_local_session_round_trip(hostname):
    auth.reset_rate_limiter()
    client = TestClient(main.app, base_url=f"http://{hostname}:8000")
    response = client.post("/login", json={"username": "patricia", "password": "demo-admin"})
    assert response.status_code == 200
    assert (
        "HttpOnly" in response.headers["set-cookie"]
        and "SameSite=lax" in response.headers["set-cookie"]
    )
    assert client.get("/app/").status_code == 200
    assert client.get("/me").json()["username"] == "patricia"
    assert client.post("/logout").status_code == 200
    assert client.get("/me").status_code == 401


def test_admin_manual_role_consistent_query_compare():
    auth.reset_rate_limiter()
    client = TestClient(main.app)
    assert (
        client.post("/login", json={"username": "patricia", "password": "demo-admin"}).status_code
        == 200
    )
    body = {
        "query": "annual recurring revenue growth",
        "role": "analyst",
        "policy_name": "full_policy",
    }
    single = client.post("/query", json=body).json()
    compare_body = {k: v for k, v in body.items() if k != "policy_name"}
    compared = client.post("/compare", json={**compare_body, "policies": ["full_policy"]}).json()[
        "results"
    ]["full_policy"]
    for result in [single, compared]:
        assert result["execution"]["identity"] == "patricia"
        assert result["execution"]["authenticated_role"] == "partner"
        assert result["execution"]["effective_role"] == "analyst"
        assert result["execution"]["simulated"] is True
        assert result["decision_trace"]["user_context"]["role"] == "analyst"
    assert single["context"] == compared["context"]


def test_audit_is_workspace_scoped_and_reports_custom_benchmark_absence(local_corpus):
    with main._session_audit_lock:
        main._session_audit[:] = [
            {"id": "q013", "workspace": "elsewhere", "user": None, "query": "foreign"},
            {"id": "q001", "workspace": "own", "user": "private", "query": "secret"},
        ]
    client = TestClient(main.app)
    response = client.get("/session-audit?workspace=own").json()
    assert response["entries"] == []
    assert response["workspace"] == "own" and response["benchmark_count"] == 0
    assert response["scope"] == "workspace_public_lab"
    with main._session_audit_lock:
        main._session_audit.clear()


def test_admin_cannot_simulate_above_actual_document_role(local_corpus):
    Path(local_corpus.users_path).write_text(
        json.dumps(
            {
                "users": [
                    {
                        "username": "admin",
                        "name": "Admin",
                        "role": "viewer",
                        "is_admin": True,
                        "password_sha256": auth.hash_password("secret"),
                    }
                ]
            }
        )
    )
    auth.reset_rate_limiter()
    client = TestClient(main.app)
    assert (
        client.post("/login", json={"username": "admin", "password": "secret"}).status_code == 200
    )
    for endpoint, extras in [("/query", {}), ("/compare", {"policies": ["full_policy"]})]:
        result = client.post(endpoint, json={"query": "annual revenue", "role": "owner", **extras})
        assert result.status_code == 403
    result = client.post("/query", json={"query": "annual revenue", "policy_name": "naive_top_k"})
    assert result.status_code == 403


def test_document_browser_enforces_authenticated_role():
    auth.reset_rate_limiter()
    client = TestClient(main.app)
    assert (
        client.post("/login", json={"username": "julia", "password": "demo-analyst"}).status_code
        == 200
    )
    response = client.get("/documents?workspace=pe-deal").json()
    assert response["visible_docs"] < response["total_docs"]
    assert {doc["min_role"] for doc in response["documents"]} == {"analyst"}
    assert client.get("/documents?workspace=saas-internal").status_code == 403


def test_ask_budget_parameter_is_part_of_cache_key(local_corpus, fake_model):
    client = TestClient(main.app)
    body = {"query": "annual revenue", "role": "viewer", "policy_name": "full_policy"}
    first = client.post("/ask", json=body).json()
    limited = client.post("/ask", json={**body, "token_budget": 0}).json()
    assert first["context"] and limited["context"] == []
    assert limited["cached"] is False and limited["model"] == "not_called"
    assert limited["decision_trace"]["result_status"] == "budget_exhausted"


def test_product_budget_cannot_exceed_policy_cap():
    auth.reset_rate_limiter()
    client = TestClient(main.app)
    client.post("/login", json={"username": "julia", "password": "demo-analyst"})
    assert (
        client.post("/query", json={"query": "annual revenue", "token_budget": 2049}).status_code
        == 400
    )
    response = client.post("/query", json={"query": "annual revenue", "token_budget": 0})
    assert response.status_code == 200 and response.json()["context"] == []


def test_transaction_setup_failure_does_not_disable_next_backup(tmp_path, monkeypatch):
    path = tmp_path / "data"
    path.write_text("old")
    original = Path.mkdir

    def fail_staging(self, *args, **kwargs):
        if self.name == ".querytrace-transaction-preparing":
            raise OSError("injected setup failure")
        return original(self, *args, **kwargs)

    with corpus_lock(str(tmp_path)):
        with monkeypatch.context() as patch:
            patch.setattr(Path, "mkdir", fail_staging)
            with pytest.raises(OSError, match="setup failure"):
                with corpus_transaction(str(tmp_path), [str(path)]):
                    pytest.fail("must fail before mutation")
        with pytest.raises(RuntimeError):
            with corpus_transaction(str(tmp_path), [str(path)]):
                path.write_text("partial")
                raise RuntimeError("rollback")
    assert path.read_text() == "old"
